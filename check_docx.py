# -*- coding: utf-8 -*-
"""交付前自检：dump 文本、查残留标记、统计图片与表格。

用法：uv run --no-project --with python-docx python -B check_docx.py 文件.docx
"""
import sys

from docx import Document
from docx.oxml.ns import qn

BAD = ["$", "**", "!!", "\\frac", "\\int", "\\dfrac"]


def main(path):
    doc = Document(path)
    bad = []
    body = doc.element.body
    imgs = len(list(body.iter(qn("a:blip"))))
    for i, p in enumerate(doc.paragraphs):
        t = p.text
        for k in BAD:
            if k in t:
                bad.append(("段落%d" % i, t[:80]))
                break
    for ti, tbl in enumerate(doc.tables):
        for ri, row in enumerate(tbl.rows):
            for ci, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    for k in BAD:
                        if k in p.text:
                            bad.append(("表%d(%d,%d)" % (ti, ri, ci), p.text[:80]))
                            break
    print("文件:", path)
    print("表格:", len(doc.tables), " 公式图片:", imgs, " 段落:", len(doc.paragraphs))
    print("残留标记:", len(bad))
    for where, txt in bad[:40]:
        print("  ", where, "|", txt)
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "数学小抄_数二背诵版.docx"))
