# -*- coding: utf-8 -*-
"""数学小抄公共样式库（供 build_数学小抄_*.py 复用）。

与政治小抄的 xc_lib 同源，差别：
1. 文本里 $...$ 会换成 LaTeX 渲染的公式图（MathBank）；
2. 单元格宽度不够时公式图自动等比缩小；
3. 表格默认字号 9pt、行距更紧，方便一页塞更多公式。
"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SONG, HEI, MONO = "宋体", "黑体", "Consolas"
RED = RGBColor(0xC0, 0x00, 0x00)
DGRAY = RGBColor(0x59, 0x59, 0x59)
DBLUE = RGBColor(0x1F, 0x4E, 0x79)
GREEN = RGBColor(0x2E, 0x6B, 0x3E)
HDR_FILL, LBL_FILL, HL_FILL = "DEEAF6", "F2F2F2", "FFF2CC"
BORDER = "BFBFBF"

TBLPR_ORDER = ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
               "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
               "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"]

BANK = None            # mathimg.MathBank，由 build 脚本注入
IMG_ALLOW = True
TOKEN = re.compile(r"(\*\*.+?\*\*|!!.+?!!|\$[^$]+\$)")


def tblpr_insert(tblPr, element):
    tag = element.tag.split("}")[1]
    idx = TBLPR_ORDER.index("w:" + tag)
    for child in tblPr:
        ctag = child.tag.split("}")[1]
        if ("w:" + ctag) in TBLPR_ORDER and TBLPR_ORDER.index("w:" + ctag) > idx:
            child.addprevious(element)
            return
    tblPr.append(element)


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def set_run(run, name=SONG, size=9, bold=None, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), name)


def add_runs(p, text, size=9, name=SONG, bold=False, color=None, img_max=None):
    """**粗体** / !!红字!! / $公式$ 混排；粗体与红字里可以再嵌套公式。img_max：公式图最大宽度(cm)。"""
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            add_runs(p, tok[2:-2], size, name, True, color, img_max)
            continue
        if tok.startswith("!!") and tok.endswith("!!") and len(tok) > 4:
            add_runs(p, tok[2:-2], size, name, True, RED, img_max)
            continue
        if tok.startswith("$") and tok.endswith("$") and len(tok) > 2:
            if not IMG_ALLOW or BANK is None:
                raise RuntimeError("公式需要 MathBank：" + tok)
            path, w_cm = BANK.get(tok[1:-1], "inline")
            if img_max:
                w_cm = min(w_cm, img_max)
            r = p.add_run()
            r.add_picture(path, width=Cm(w_cm))
            continue
        r = p.add_run(tok)
        set_run(r, name, size, bold, color)


def para(doc, text="", size=9, bold=False, color=None, align=None,
         before=0, after=0, line=1.1, keep=False, name=SONG, indent=None, img_max=17.5):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    if indent is not None:
        fmt.left_indent = Cm(indent)
    if align is not None:
        p.alignment = align
    if text:
        add_runs(p, text, size, name, bold, color, img_max=img_max)
    return p


def note(doc, text, after=6):
    return para(doc, "※ " + text, size=8.5, color=DGRAY, after=after, line=1.05)


def section(doc, text, brk=False):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before, fmt.space_after = Pt(10 if not brk else 2), Pt(3)
    fmt.keep_with_next = True
    fmt.page_break_before = brk
    fmt.left_indent = Cm(0.08)
    pPr = p._element.get_or_add_pPr()
    pPr.append(el("w:shd", val="clear", color="auto", fill=HDR_FILL))
    add_runs(p, text, size=11.5, name=HEI, bold=True)
    return p


def sub(doc, text):
    return para(doc, text, size=10, name=HEI, bold=True, color=DBLUE,
                before=7, after=2.5, keep=True)


def shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(el("w:shd", val="clear", color="auto", fill=fill))


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(el("w:tblHeader", val="true"))


def configure_table(t, total_cm, borders_color=BORDER):
    tblPr = t._tbl.tblPr
    for tag in ("w:tblW", "w:tblBorders", "w:tblLayout", "w:tblCellMar"):
        for e in tblPr.findall(qn(tag)):
            tblPr.remove(e)
    tblpr_insert(tblPr, el("w:tblW", w=str(int(total_cm * 567)), type="dxa"))
    brd = OxmlElement("w:tblBorders")
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        brd.append(el("w:" + tag, val="single", sz="4", space="0", color=borders_color))
    tblpr_insert(tblPr, brd)
    tblpr_insert(tblPr, el("w:tblLayout", type="fixed"))
    mar = OxmlElement("w:tblCellMar")
    for tag, v in (("top", 30), ("left", 68), ("bottom", 30), ("right", 68)):
        mar.append(el("w:" + tag, w=str(v), type="dxa"))
    tblpr_insert(tblPr, mar)


def set_grid_widths(t, widths):
    """把列宽写进 tblGrid：WPS 认 tcW，Word/LibreOffice 认 gridCol，两边都写才一致。"""
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(round(w * 567))))


def ensure_table_gap(doc):
    """相邻两张表之间留 5pt 空隙，否则会被 Word/WPS 并成一张，跨页时表头串行。"""
    body = doc.element.body
    for child in reversed(list(body.iterchildren())):
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:tbl"):
            p = doc.add_paragraph()
            fmt = p.paragraph_format
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.line_spacing = Pt(5)
        break


def make_table(doc, widths, rows, head=None, header=None, hl=(), label_col=False, size=9,
               cant_split=True, keep_together=False, center_formula=True):
    """head：表头文字列表（全空则自动不画表头）；header 可强制指定。"""
    if header is None:
        header = bool(head) and any(h.strip() for h in head)
    if header:
        rows = [list(head)] + [list(r) for r in rows]
    ensure_table_gap(doc)
    t = doc.add_table(rows=len(rows), cols=len(widths))
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    configure_table(t, sum(widths))
    set_grid_widths(t, widths)
    for ri, rowdata in enumerate(rows):
        trPr = t.rows[ri]._tr.get_or_add_trPr()
        if cant_split:
            trPr.append(el("w:cantSplit", val="true"))
        for ci, val in enumerate(rowdata):
            cell = t.cell(ri, ci)
            cell.width = Cm(widths[ci])
            is_header = header and ri == 0
            if is_header:
                shade(cell, HDR_FILL)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            elif ri in hl:
                shade(cell, HL_FILL)
            elif label_col and ci == 0:
                shade(cell, LBL_FILL)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            elif center_formula:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            lines = str(val).split("\n")
            for li, line in enumerate(lines):
                p = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.08
                if is_header:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.keep_with_next = True
                elif center_formula and line.strip().startswith("$") and line.strip().endswith("$") and line.count("$") == 2:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if line:
                    add_runs(p, line, size, bold=is_header, img_max=widths[ci] - 0.42)
                    if keep_together and ri < len(rows) - 1:
                        p.paragraph_format.keep_with_next = True
    if header:
        set_repeat_header(t.rows[0])
    return t


def merge_v(t, col, r1, r2, fill=None):
    cell = t.cell(r1, col).merge(t.cell(r2, col))
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        shade(cell, fill)
    return cell


def add_page_field(p, size=7.5):
    run = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._element.append(f1); run._element.append(it); run._element.append(f2)
    set_run(run, SONG, size, False, DGRAY)


def init_doc(subject):
    """A4 窄边距、宋体正文、页脚“数学小抄 · <主题> ｜ 第 X 页”。"""
    doc = Document()
    doc.core_properties.title = "数学小抄 · " + subject
    doc.core_properties.author = "数学小抄"
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.35), Cm(1.15)
    sec.left_margin, sec.right_margin = Cm(1.3), Cm(1.3)
    st = doc.styles["Normal"]
    st.font.name = SONG
    st.font.size = Pt(9)
    try:
        st.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
    except Exception:
        pass
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(fp, "数学小抄 · " + subject + " ｜ 第 ", 7.5, color=DGRAY)
    add_page_field(fp)
    add_runs(fp, " 页", 7.5, color=DGRAY)
    return doc


def title_block(doc, main, sub_, meta, usage):
    para(doc, main, size=16, name=HEI, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    para(doc, sub_, size=10.5, name=HEI, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    para(doc, meta, size=8.5, color=DGRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    para(doc, usage, size=8.5, color=DGRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)


def save(doc, out):
    doc.save(out)
    print("已生成:", out)
    print("表格数:", len(doc.tables))
